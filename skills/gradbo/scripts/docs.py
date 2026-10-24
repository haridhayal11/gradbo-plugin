# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1", "pypdfium2>=4.30", "pillow>=10"]
# ///
"""Build and check application documents on any OS.

Run with `uv run docs.py <command> ...` (installs the dependencies on first use), or
`python docs.py ...` after `python -m pip install python-docx pypdfium2 pillow`.

Commands
  check FILE [--one-page]     .tex (compiled with Tectonic), .pdf or .docx: extract the text,
                              render the pages to render/, report pages, size, words and SHA-256.
                              Exits 1 when a check fails. A .docx is converted with LibreOffice
                              or Word when either is installed, so its pages can be rendered too.
  text FILE                   print the document's text (.tex is read from its compiled PDF)
  compare BASE TAILORED [--allow FILE]... [--allow-term TEXT]...
                              list numbers, names and terms in TAILORED that BASE never mentions,
                              and BASE lines that TAILORED dropped. Exits 1 when anything was added.
                              Terms in an --allow file (profile/facts.md) and each --allow-term
                              (the employer's name, the role title) count as known.
  docx-dump FILE              print every paragraph (body and tables, in order) as JSON:
                              {"i", "style", "text", "where"}; the "i" numbers are what docx-edit uses
  docx-edit FILE EDITS OUT    apply a JSON list of edits (indices from docx-dump of FILE) and save OUT:
                                {"op": "set", "i": 12, "text": "..."}     replace a paragraph's text,
                                                                         keeping its first run's formatting
                                {"op": "delete", "i": 7}
                                {"op": "move", "i": 9, "before": 4}     both in the same section or cell
  letter TEXT OUT.docx        a plain one-page cover letter from a text file: first line the name,
                              second line the contact details, then paragraphs separated by blank lines
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

LIGATURES = "ﬀﬁﬂﬃﬄﬅﬆ"
A4 = (595.3, 841.9)
LETTER_WORDS = 450  # beyond this a cover letter rarely fits one A4 page


def fail(message: str, code: int = 2):
    print(message, file=sys.stderr)
    sys.exit(code)


def need(module: str, package: str):
    try:
        return __import__(module)
    except ImportError:
        fail(f"missing Python package {package}: run this script with `uv run`, or `python -m pip install {package}`")


# -- PDF -------------------------------------------------------------------

def compile_tex(src: Path) -> Path:
    tectonic = shutil.which("tectonic")
    if not tectonic:
        fail("Tectonic is not installed (https://tectonic-typesetting.github.io). "
             "Send the candidate's own PDF unchanged and say so in the receipt.")
    result = subprocess.run([tectonic, "--keep-logs", src.name], cwd=src.parent, capture_output=True, text=True)
    if result.returncode:
        print(result.stderr[-3000:], file=sys.stderr)
        fail(f"compile failed; see {src.with_suffix('.log')}", 1)
    return src.with_suffix(".pdf")


def docx_to_pdf(src: Path) -> Path | None:
    """Convert with LibreOffice, or with Word through docx2pdf, when either is available."""
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        out = Path(tempfile.mkdtemp())
        subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(src)],
                       capture_output=True, timeout=180)
        pdf = out / (src.stem + ".pdf")
        return pdf if pdf.exists() else None
    try:
        from docx2pdf import convert  # Word on Windows and macOS
    except ImportError:
        return None
    out = Path(tempfile.mkdtemp()) / (src.stem + ".pdf")
    try:
        convert(str(src), str(out))
    except Exception:
        return None
    return out if out.exists() else None


def pdf_text(pdf: Path) -> tuple[str, list]:
    pdfium = need("pypdfium2", "pypdfium2")
    doc = pdfium.PdfDocument(str(pdf))
    pages = [doc[i] for i in range(len(doc))]
    text = "\n\f".join(p.get_textpage().get_text_range() for p in pages)
    return text.replace("\r\n", "\n"), pages


def check_pdf(pdf: Path, name: str, out_dir: Path, one_page: bool) -> int:
    text, pages = pdf_text(pdf)
    (out_dir / f"{name}.txt").write_text(text, encoding="utf-8")
    render = out_dir / "render"
    render.mkdir(exist_ok=True)
    for old in render.glob(f"{name}-page-*.png"):
        old.unlink()
    images = []
    for n, page in enumerate(pages, 1):
        image = render / f"{name}-page-{n}.png"
        page.render(scale=150 / 72).to_pil().save(image)
        images.append(str(image))
    width, height = pages[0].get_size() if pages else (0, 0)
    words = len(text.split())
    print(f"pdf: {pdf}")
    print(f"pages: {len(pages)}")
    print(f"page size: {width:.0f} x {height:.0f} pt{' (A4)' if abs(width - A4[0]) < 2 and abs(height - A4[1]) < 2 else ''}")
    print(f"extracted words: {words}")
    print(f"text: {out_dir / (name + '.txt')}")
    print("render: " + " ".join(images))
    print(f"sha256: {hashlib.sha256(pdf.read_bytes()).hexdigest()}")
    status = 0
    if words < 30:
        print("FAIL: almost no text extracted; an applicant tracking system would read nothing", file=sys.stderr)
        status = 1
    if any(ch in text for ch in LIGATURES):
        print("FAIL: ligature characters (ﬀ ﬁ ﬂ) in the text; turn ligatures off", file=sys.stderr)
        status = 1
    if one_page and len(pages) != 1:
        print(f"FAIL: {len(pages)} pages; a cover letter must fit on one A4 page", file=sys.stderr)
        status = 1
    print("Now look at every rendered page for clipping, overlap and text running off the page.")
    return status


# -- DOCX ------------------------------------------------------------------

def docx_paragraphs(document):
    """Every paragraph in document order, including those inside tables (CVs often use them)."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    def walk(container, where):
        element = container.element.body if hasattr(container, "element") and hasattr(container.element, "body") else container._element
        for child in element.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                yield Paragraph(child, container), where
            elif tag == "tbl":
                table = Table(child, container)
                for r, row in enumerate(table.rows):
                    for c, cell in enumerate(row.cells):
                        yield from walk(cell, f"{where}table r{r} c{c} ")
    yield from walk(document, "")


def docx_text(path: Path) -> str:
    docx = need("docx", "python-docx")
    return "\n".join(p.text for p, _ in docx_paragraphs(docx.Document(str(path))))


def check_docx(path: Path, one_page: bool) -> int:
    pdf = docx_to_pdf(path)
    if pdf:
        return check_pdf(pdf, path.stem, path.parent, one_page)
    text = docx_text(path)
    (path.parent / f"{path.stem}.txt").write_text(text, encoding="utf-8")
    words = len(text.split())
    print(f"docx: {path}")
    print(f"extracted words: {words}")
    print(f"sha256: {hashlib.sha256(path.read_bytes()).hexdigest()}")
    print("render: unavailable (install LibreOffice to render Word files); the layout was not checked")
    status = 0
    if words < 30:
        print("FAIL: almost no text in the document", file=sys.stderr)
        status = 1
    if one_page and words > LETTER_WORDS:
        print(f"FAIL: {words} words; a one-page cover letter stays under about {LETTER_WORDS}", file=sys.stderr)
        status = 1
    return status


def docx_dump(path: Path):
    docx = need("docx", "python-docx")
    rows = [{"i": i, "style": p.style.name if p.style is not None else "", "text": p.text, "where": where.strip() or "body"}
            for i, (p, where) in enumerate(docx_paragraphs(docx.Document(str(path))))]
    print(json.dumps(rows, indent=1, ensure_ascii=False))


def set_text(paragraph, text: str):
    """Replace the text, keeping formatting: leading runs whose text is unchanged (a bold
    "Label:" for instance) stay as they are, and the rest goes into the next run's format."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    kept, rest = 0, text
    while kept < len(runs) - 1 and runs[kept].text and rest.startswith(runs[kept].text):
        rest = rest[len(runs[kept].text):]
        kept += 1
    runs[kept].text = rest
    for run in runs[kept + 1:]:
        run._r.getparent().remove(run._r)


def docx_edit(path: Path, edits_path: Path, out: Path):
    docx = need("docx", "python-docx")
    document = docx.Document(str(path))
    paragraphs = [p for p, _ in docx_paragraphs(document)]
    edits = json.loads(edits_path.read_text(encoding="utf-8"))
    for edit in edits:
        if edit.get("op") == "set":
            set_text(paragraphs[edit["i"]], edit["text"])
    for edit in edits:
        if edit.get("op") == "move":
            moving, anchor = paragraphs[edit["i"]]._p, paragraphs[edit["before"]]._p
            if moving.getparent() is not anchor.getparent():
                fail(f"move {edit['i']} before {edit['before']}: they are in different sections or cells", 1)
            anchor.addprevious(moving)
    for edit in sorted((e for e in edits if e.get("op") == "delete"), key=lambda e: -e["i"]):
        element = paragraphs[edit["i"]]._p
        element.getparent().remove(element)
    unknown = [e for e in edits if e.get("op") not in ("set", "move", "delete")]
    if unknown:
        fail(f"unknown edit: {unknown[0]}", 1)
    document.save(str(out))
    print(f"saved {out} ({len(edits)} edits)")


def letter(text_path: Path, out: Path):
    docx = need("docx", "python-docx")
    from docx.shared import Mm, Pt

    blocks = [b.strip() for b in re.split(r"\n\s*\n", text_path.read_text(encoding="utf-8").strip()) if b.strip()]
    head = blocks.pop(0).splitlines() if blocks else []
    document = docx.Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(25))
    style = document.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    style.paragraph_format.space_after = Pt(8)
    if head:
        run = document.add_paragraph().add_run(head[0])
        run.bold, run.font.size = True, Pt(14)
        for line in head[1:]:
            document.add_paragraph(line)
    for block in blocks:
        document.add_paragraph(" ".join(line.strip() for line in block.splitlines()))
    document.save(str(out))
    print(f"saved {out}")


# -- comparing a tailored CV with its base ---------------------------------

TOKEN = re.compile(r"[£$€]?\d[\d.,]*%?|[A-Za-z][A-Za-z0-9+#./-]*[A-Za-z0-9+#]|[A-Za-z]")


def text_of(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return docx_text(path)
    if suffix == ".pdf":
        return pdf_text(path)[0]
    if suffix == ".tex":
        pdf = path.with_suffix(".pdf")
        if not pdf.exists() or pdf.stat().st_mtime < path.stat().st_mtime:
            pdf = compile_tex(path)
        return pdf_text(pdf)[0]
    return path.read_text(encoding="utf-8")


def notable(token: str) -> bool:
    """Numbers, and words that carry a claim: capitalised names, acronyms, tool names."""
    return bool(re.match(r"[£$€]?\d", token)) or token[0].isupper() or bool(re.search(r"[+#./\d]", token))


def compare(base_path: Path, tailored_path: Path, allowed: list[str] | None = None) -> int:
    base, tailored = text_of(base_path), text_of(tailored_path)
    known = {t.lower().strip(".,") for t in TOKEN.findall(" ".join([base, *(allowed or [])]))}
    added = sorted({t.strip(".,") for t in TOKEN.findall(tailored) if notable(t) and t.lower().strip(".,") not in known})
    norm = lambda line: re.sub(r"\W+", " ", line).strip().lower()
    tailored_lines = {norm(l) for l in tailored.splitlines()}
    tailored_words = set(norm(tailored).split())
    dropped = [l.strip() for l in base.splitlines()
               if len(norm(l).split()) >= 4 and norm(l) not in tailored_lines
               and len(set(norm(l).split()) - tailored_words) > len(norm(l).split()) // 2]
    if added:
        print("ADDED (not in the base CV; remove each unless the candidate's facts support it):")
        for token in added:
            print(f"  {token}")
    if dropped:
        print("DROPPED (in the base CV, missing from the tailored one; fine if deliberate):")
        for line in dropped:
            print(f"  {line}")
    if not added and not dropped:
        print("No new names, numbers or terms, and nothing dropped.")
    return 1 if added else 0


def main(argv: list[str]) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    if not argv or argv[0] in ("-h", "--help"):
        print(__doc__)
        return 0
    command, args = argv[0], argv[1:]
    one_page = "--one-page" in args
    args = [a for a in args if a != "--one-page"]
    allowed: list[str] = []
    while any(a in ("--allow", "--allow-term") for a in args):
        i = next(i for i, a in enumerate(args) if a in ("--allow", "--allow-term"))
        if i + 1 >= len(args):
            fail(f"{args[i]} needs a value")
        flag, value = args[i], args[i + 1]
        allowed.append(Path(value).read_text(encoding="utf-8") if flag == "--allow" else value)
        del args[i:i + 2]
    if command == "check" and len(args) == 1:
        path = Path(args[0]).resolve()
        suffix = path.suffix.lower()
        if suffix == ".tex":
            return check_pdf(compile_tex(path), path.stem, path.parent, one_page)
        if suffix == ".pdf":
            return check_pdf(path, path.stem, path.parent, one_page)
        if suffix == ".docx":
            return check_docx(path, one_page)
        fail("check takes a .tex, .pdf or .docx file")
    if command == "text" and len(args) == 1:
        print(text_of(Path(args[0]).resolve()))
        return 0
    if command == "compare" and len(args) == 2:
        return compare(Path(args[0]).resolve(), Path(args[1]).resolve(), allowed)
    if command == "docx-dump" and len(args) == 1:
        docx_dump(Path(args[0]).resolve())
        return 0
    if command == "docx-edit" and len(args) == 3:
        docx_edit(*(Path(a).resolve() for a in args))
        return 0
    if command == "letter" and len(args) == 2:
        letter(Path(args[0]).resolve(), Path(args[1]).resolve())
        return 0
    fail(f"usage: see `docs.py --help`")


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
